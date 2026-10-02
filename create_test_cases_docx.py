import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx(output_path):
    doc = Document()
    
    # Page Setup - Standard Margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("BlogSpace — System Test Cases Report")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
    title_p.paragraph_format.space_after = Pt(4)

    subtitle_p = doc.add_paragraph()
    sub_run = subtitle_p.add_run("Chapter 6: System Testing & Quality Assurance (Test Suite T01 - T14)")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    subtitle_p.paragraph_format.space_after = Pt(18)

    # Summary Metadata Box
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Project:", " BlogSpace Platform"), ("Date:", " October 2, 2026"), ("Tested By:", " QA Team"), ("Environment:", " Localhost / Node.js")],
        [("Total Cases:", " 14"), ("Passed:", " 14"), ("Failed:", " 0"), ("Pass Rate:", " 100%")]
    ]
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            lbl, val = meta_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            r1 = p.add_run(lbl)
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = RGBColor(15, 23, 42)
            r2 = p.add_run(val)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section Title
    sec_p = doc.add_paragraph()
    sec_run = sec_p.add_run("Table 6.1: Functional Test Execution Matrix")
    sec_run.font.name = "Calibri"
    sec_run.font.size = Pt(14)
    sec_run.font.bold = True
    sec_run.font.color.rgb = RGBColor(15, 23, 42)
    sec_p.paragraph_format.space_after = Pt(8)

    # Main Test Cases Table
    test_cases = [
        ("T01", "Register with valid data", "Name, unique e-mail, password", "Account created; success message", "Pass"),
        ("T02", "Register with existing e-mail", "E-mail already in database", "Error: e-mail already registered", "Pass"),
        ("T03", "Register with empty field", "Leave password blank", "Validation error shown", "Pass"),
        ("T04", "Login with valid credentials", "Correct e-mail and password", "Token issued; redirected to home", "Pass"),
        ("T05", "Login with wrong password", "Correct e-mail, wrong password", "Error: invalid credentials", "Pass"),
        ("T06", "Login as blocked user", "Blocked account", "Access denied message", "Pass"),
        ("T07", "Create post as author", "Valid title, content, category", "Post saved; visible in feed", "Pass"),
        ("T08", "Create post without title", "Empty title", "Validation error; nothing saved", "Pass"),
        ("T09", "Save post as draft", "Click Save Draft", "Post stored with status draft; not public", "Pass"),
        ("T10", "Edit own post", "Change title and save", "Changes saved and shown", "Pass"),
        ("T11", "Edit another author's post", "PUT request for other's post", "403 Forbidden error returned", "Pass"),
        ("T12", "Delete own post", "Click Delete and confirm", "Post removed from list", "Pass"),
        ("T13", "Upload cover image", "JPG under size limit", "Image stored and displayed", "Pass"),
        ("T14", "Upload invalid file", "PDF or oversize image", "Upload rejected with message", "Pass"),
    ]

    table = doc.add_table(rows=len(test_cases) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Column Widths
    widths = [Inches(0.6), Inches(2.1), Inches(1.8), Inches(2.1), Inches(0.8)]
    headers = ["ID", "Test Case", "Input / Action", "Expected Result", "Result"]

    # Header Row Styling
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "4F46E5") # Primary Indigo
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 4] else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, data in enumerate(test_cases, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "FFFFFF" if row_idx % 2 != 0 else "F8FAFC"
        
        for col_idx, text in enumerate(data):
            row_cells[col_idx].width = widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=90, bottom=90, left=100, right=100)
            
            p = row_cells[col_idx].paragraphs[0]
            if col_idx in [0, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(79, 70, 229)
            elif col_idx == 4:
                run.font.bold = True
                run.font.color.rgb = RGBColor(16, 185, 129) # Success Green
            else:
                run.font.color.rgb = RGBColor(30, 41, 59)

    doc.save(output_path)
    print(f"Successfully generated DOCX file at: {output_path}")

if __name__ == '__main__':
    target_path = r"C:\Users\ACER\.gemini\antigravity\scratch\blogspace\BlogSpace_System_Test_Cases.docx"
    brain_path = r"C:\Users\ACER\.gemini\antigravity\brain\2f272290-eebc-4be7-904e-588f7f6e4954\BlogSpace_System_Test_Cases.docx"
    create_docx(target_path)
    create_docx(brain_path)
