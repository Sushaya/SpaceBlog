import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_validation_rules_docx(output_path):
    doc = Document()
    
    # Page Setup - Standard Margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("BlogSpace — System Input Validation Specifications")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
    title_p.paragraph_format.space_after = Pt(4)

    subtitle_p = doc.add_paragraph()
    sub_run = subtitle_p.add_run("Chapter 6: Data Integrity & Validation Enforcement Specification")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    subtitle_p.paragraph_format.space_after = Pt(16)

    # Section Title
    sec_p = doc.add_paragraph()
    sec_run = sec_p.add_run("Table 6.4: Input Validation Rules")
    sec_run.font.name = "Calibri"
    sec_run.font.size = Pt(14)
    sec_run.font.bold = True
    sec_run.font.color.rgb = RGBColor(15, 23, 42)
    sec_p.paragraph_format.space_after = Pt(8)

    # Table 6.4 Data
    validation_data = [
        ("Name", "Required; 2 to 50 characters", "Please enter a valid name"),
        ("E-mail", "Required; valid e-mail format; unique", "Enter a valid e-mail / already registered"),
        ("Password", "Required; at least 8 characters", "Password must be at least 8 characters"),
        ("Post title", "Required; 5 to 150 characters", "Title is required (5-150 characters)"),
        ("Post content", "Required; minimum 100 characters", "Content is too short"),
        ("Cover image", "JPG, PNG or WEBP; maximum 2 MB", "Invalid image type or size"),
        ("Comment", "Required; maximum 500 characters", "Comment cannot be empty or too long")
    ]

    table = doc.add_table(rows=len(validation_data) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Column Widths
    widths = [Inches(1.5), Inches(2.7), Inches(2.8)]
    headers = ["Field", "Rule", "Message shown"]

    # Header Row Styling
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "4F46E5") # Primary Indigo
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, data in enumerate(validation_data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "FFFFFF" if row_idx % 2 != 0 else "F8FAFC"
        
        for col_idx, text in enumerate(data):
            row_cells[col_idx].width = widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=90, bottom=90, left=100, right=100)
            
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(79, 70, 229)
            elif col_idx == 2:
                run.font.color.rgb = RGBColor(220, 38, 38) # Crimson Red for messages
            else:
                run.font.color.rgb = RGBColor(30, 41, 59)

    doc.save(output_path)
    print(f"Successfully generated Table 6.4 DOCX file at: {output_path}")

if __name__ == '__main__':
    target_path = r"C:\Users\ACER\.gemini\antigravity\scratch\blogspace\BlogSpace_Table_6_4_Input_Validation_Rules.docx"
    create_validation_rules_docx(target_path)
